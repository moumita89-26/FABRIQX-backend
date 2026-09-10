import csv
from decimal import Decimal, InvalidOperation
from io import TextIOWrapper
from pathlib import Path
from zipfile import BadZipFile
from openpyxl.utils.exceptions import InvalidFileException

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook

from . import models as m


SAMPLE_COLUMNS = (
    "category", "product_name", "brand", "short_description", "description",
    "regular_price", "sale_price", "sale", "is_new_arrival", "sku", "size",
    "color", "color_code", "stock_quantity",
)

SAMPLE_ROW = {
    "category": "Dresses", "product_name": "Silk Celebration Dress", "brand": "FABRIQX",
    "short_description": "Premium silk dress",
    "description": "Detailed product description", "regular_price": "2999.00", "sale_price": "2499.00",
    "sale": "true", "is_new_arrival": "true",
    "sku": "SILK-DRESS-M-RED", "size": "M", "color": "Red", "color_code": "#9B1C1C",
    "stock_quantity": "25",
}


def _rows(upload):
    suffix = Path(upload.name).suffix.lower()
    if suffix == ".csv":
        yield from csv.DictReader(TextIOWrapper(upload.file, encoding="utf-8-sig"))
        return
    if suffix == ".xlsx":
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as exc:
            raise ValueError("Invalid Excel file. Upload a valid .xlsx workbook using the product import template.") from exc
        try:
            sheet = workbook.active
            if sheet is None:
                raise ValueError("The Excel workbook has no worksheet.")
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(iterator, ())]
            required = ("category", "product_name", "regular_price", "sku")
            missing = [name for name in required if name not in headers]
            if missing:
                raise ValueError("Missing required columns: " + ", ".join(missing) + ". Use the product import template.")
            for values in iterator:
                if any(value is not None and str(value).strip() for value in values):
                    yield dict(zip(headers, values))
        finally:
            workbook.close()
        return
    raise ValueError("Upload a .csv or .xlsx file.")


def _bool(value):
    return str(value or "").strip().lower() in {"1", "true", "yes", "y", "on"}


def _decimal(value, default=None):
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError(f"Invalid decimal value: {value}") from exc


def _integer(value, default=0):
    if value in (None, ""):
        return default
    return int(value)


def _format_error(exc):
    if isinstance(exc, ValidationError):
        if hasattr(exc, "message_dict"):
            messages = []
            for field, field_errors in exc.message_dict.items():
                label = "Product" if field == "__all__" else field.replace("_", " ").capitalize()
                messages.append(f"{label}: {'; '.join(field_errors)}")
            return " ".join(messages)
        return "; ".join(exc.messages)
    return str(exc)


def _row_list(row_numbers):
    ranges = []
    start = previous = row_numbers[0]
    for number in row_numbers[1:]:
        if number == previous + 1:
            previous = number
            continue
        ranges.append(str(start) if start == previous else f"{start}-{previous}")
        start = previous = number
    ranges.append(str(start) if start == previous else f"{start}-{previous}")
    return ", ".join(ranges)


def _process_row(raw):
    row = {str(key).strip(): value for key, value in raw.items() if key}
    required = ("category", "product_name", "regular_price", "sku")
    missing = [name for name in required if not str(row.get(name, "")).strip()]
    if missing:
        raise ValueError(f"Missing required column values: {', '.join(missing)}")

    category, _ = m.Category.objects.get_or_create(
        name=str(row["category"]).strip(), parent=None,
        defaults={"is_active": True},
    )

    slug = slugify(str(row["product_name"]))
    product_defaults = {
        "category": category,
        "name": str(row["product_name"]).strip(),
        "brand": str(row.get("brand") or "").strip(),
        "short_description": str(row.get("short_description") or "").strip(),
        "description": str(row.get("description") or "").strip(),
        "regular_price": _decimal(row["regular_price"], Decimal("0")),
        "sale_price": _decimal(row.get("sale_price")),
        "is_trending": _bool(row.get("sale")),
        "is_new_arrival": _bool(row.get("is_new_arrival")),
    }
    product = m.Product.objects.filter(slug=slug).first()
    created_product = product is None
    if product:
        for key, value in product_defaults.items():
            setattr(product, key, value)
    else:
        product = m.Product(slug=slug, **product_defaults)
    product.full_clean()
    product.save()

    variant_defaults = {
        "product": product,
        "size": str(row.get("size") or "").strip(),
        "color": str(row.get("color") or "").strip(),
        "color_code": str(row.get("color_code") or "").strip(),
        "stock_quantity": _integer(row.get("stock_quantity")),
        "is_active": True,
    }
    variant, variant_created = m.ProductVariant.objects.update_or_create(sku=str(row["sku"]).strip(), defaults=variant_defaults)

    return created_product, variant_created, False


def import_products(upload, batch_size=500):
    counts = {"rows": 0, "succeeded": 0, "products": 0, "variants": 0, "images": 0, "failed": 0}
    error_rows = {}
    batch = []

    def process_batch(rows):
        for row_number, row in rows:
            try:
                with transaction.atomic():
                    product_created, variant_created, image_created = _process_row(row)
                counts["products"] += int(product_created)
                counts["variants"] += int(variant_created)
                counts["images"] += int(image_created)
                counts["succeeded"] += 1
            except Exception as exc:
                counts["failed"] += 1
                product_name = str(row.get("product_name") or "").strip()
                sku = str(row.get("sku") or "").strip()
                identifier = " / ".join(value for value in (product_name, sku) if value)
                prefix = f"{identifier}: " if identifier else ""
                error = prefix + _format_error(exc)
                if error in error_rows:
                    error_rows[error].append(row_number)
                elif len(error_rows) < 50:
                    error_rows[error] = [row_number]

    for row_number, row in enumerate(_rows(upload), start=2):
        counts["rows"] += 1
        batch.append((row_number, row))
        if len(batch) >= batch_size:
            process_batch(batch)
            batch.clear()
    process_batch(batch)
    errors = [
        f"{'Row' if len(rows) == 1 else 'Rows'} {_row_list(rows)}: {error}"
        for error, rows in error_rows.items()
    ]
    return counts, errors
